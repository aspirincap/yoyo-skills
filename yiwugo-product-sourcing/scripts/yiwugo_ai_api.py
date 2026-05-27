#!/usr/bin/env python3
"""No-browser client for Yiwugo AI product sourcing.

Flow:
1. Stream AI answers from aiapi.yiwugo.com/forwebChat.
2. Extract product IDs from final answers.
3. Fetch product cards from www.yiwugo.com/api/product/aiGuideProducts.htm.
4. Normalize, de-duplicate, score, and write Excel plus HTML reports.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import io
import json
import re
import sys
import urllib.parse
import urllib.request
import uuid
from pathlib import Path
from typing import Any

WEB_ORIGIN = "https://www.yiwugo.com"
AI_CHAT_URL = "https://aiapi.yiwugo.com/forwebChat"
AI_PAGE_URL = f"{WEB_ORIGIN}/ai.html"
PRODUCT_API_URL = f"{WEB_ORIGIN}/api/product/aiGuideProducts.htm"
BROWSER_UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/147.0.0.0 Safari/537.36"

RISK_KEYWORDS = {
    "electrical": ["led", "usb", "type-c", "充电", "电池", "太阳能", "照明", "电子", "电动", "露营灯"],
    "ip": ["同款", "联名", "迪士尼", "三丽鸥", "白蛇", "青蛇", "动漫", "卡通", "ip"],
    "fragile": ["玻璃", "琉璃", "陶瓷", "水晶", "珍珠", "流苏"],
    "children": ["儿童", "婴儿", "宝宝", "幼儿", "玩具"],
    "food_contact": ["餐具", "水杯", "水壶", "饭盒", "食品", "硅胶杯"],
    "seasonal": ["圣诞", "万圣", "情人节", "春节", "秋冬", "节日"],
}


def request_text(url: str, headers: dict[str, str] | None = None, timeout: int = 30) -> str:
    req = urllib.request.Request(url, headers=headers or {})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read().decode("utf-8", "replace")


def request_bytes(url: str, headers: dict[str, str] | None = None, timeout: int = 30) -> bytes:
    req = urllib.request.Request(url, headers=headers or {})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return resp.read()


def get_csrf_token() -> str:
    html_text = request_text(AI_PAGE_URL, {"User-Agent": BROWSER_UA, "Accept": "text/html"})
    match = re.search(r'"csrf"\s*:\s*"([^"]+)"', html_text)
    if not match:
        raise RuntimeError("Could not find csrf token in ai.html")
    return match.group(1)


def extract_product_ids(text: str) -> list[str]:
    ids: list[str] = []
    seen: set[str] = set()
    for match in re.finditer(r"<product>(\d+)</product>|\b(9\d{8})\b", text):
        product_id = match.group(1) or match.group(2)
        if product_id and product_id not in seen:
            ids.append(product_id)
            seen.add(product_id)
    return ids


def stream_chat(message: str, client_id: str | None = None, session_id: str | None = None, timeout: int = 120) -> dict[str, Any]:
    client_id = client_id or str(uuid.uuid4())
    params = {"message": message, "userAgent": client_id}
    if session_id:
        params["sessionId"] = session_id
    url = f"{AI_CHAT_URL}?{urllib.parse.urlencode(params)}"
    req = urllib.request.Request(
        url,
        headers={
            "Accept": "text/event-stream",
            "Cache-Control": "no-cache",
            "Referer": f"{WEB_ORIGIN}/",
            "User-Agent": BROWSER_UA,
        },
    )

    events: list[str] = []
    router_chunks: list[str] = []
    answer_chunks: list[str] = []
    final_answer: str | None = None
    created_session_id: str | None = None
    metadata: list[Any] = []

    with urllib.request.urlopen(req, timeout=timeout) as resp:
        while True:
            line = resp.readline()
            if not line:
                break
            text = line.decode("utf-8", "replace").strip()
            if not text or text.startswith(":") or not text.startswith("data:"):
                continue
            try:
                payload = json.loads(text[5:].strip())
            except json.JSONDecodeError:
                continue

            event = payload.get("event")
            data = payload.get("data")
            if event:
                events.append(event)
            if event == "session_created" and isinstance(data, dict):
                created_session_id = data.get("session_id")
            elif event == "metadata":
                metadata.append(data)
            elif isinstance(data, str):
                if event and event.startswith("router_"):
                    router_chunks.append(data)
                elif event == "final_answer":
                    final_answer = data
                elif event and event.startswith("finalizer_"):
                    answer_chunks.append(data)

    answer = final_answer if final_answer is not None else "".join(answer_chunks)
    return {
        "client_id": client_id,
        "session_id": created_session_id or session_id,
        "events": sorted(set(events)),
        "router_text": "".join(router_chunks),
        "answer": answer,
        "product_ids": extract_product_ids(answer),
        "metadata": metadata,
    }


def fetch_products(product_ids: list[str]) -> list[dict[str, Any]]:
    if not product_ids:
        return []
    csrf = get_csrf_token()
    ids = ",".join(product_ids)
    url = f"{PRODUCT_API_URL}?{urllib.parse.urlencode({'ids': ids})}"
    text = request_text(
        url,
        {
            "Accept": "application/json, text/plain, */*",
            "Referer": AI_PAGE_URL,
            "X-Requested-With": "XMLHttpRequest",
            "x-csrf-token": csrf,
            "User-Agent": BROWSER_UA,
        },
    )
    payload = json.loads(text)
    if str(payload.get("code")) != "1":
        raise RuntimeError(f"Product API failed: {payload}")
    by_id = {str(item.get("id")): item for item in payload.get("content", {}).get("data", {}).get("productIds", [])}
    return [by_id[pid] for pid in product_ids if pid in by_id]


def cents_to_float(value: Any) -> float | None:
    try:
        ivalue = int(value)
    except (TypeError, ValueError):
        return None
    if ivalue <= 0:
        return None
    return ivalue / 100


def numeric_sort_value(value: Any, default: float = 10**9) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def format_yuan(value: float | None) -> str:
    return "" if value is None else f"¥{value:.2f}"


def normalize_url(value: Any) -> str:
    url = str(value or "").strip()
    if not url:
        return ""
    if url.startswith("//"):
        return f"https:{url}"
    if url.startswith("/"):
        return f"{WEB_ORIGIN}{url}"
    return url


def price_display(item: dict[str, Any]) -> str:
    low = cents_to_float(item.get("sellPrice"))
    high = cents_to_float(item.get("maxPrice"))
    if low is None and high is None:
        return "联系商家"
    if high is None or high == low:
        return format_yuan(low)
    return f"{format_yuan(low)}-{format_yuan(high)}"


def risk_flags(title: str) -> list[str]:
    lowered = title.lower()
    flags: list[str] = []
    for flag, keywords in RISK_KEYWORDS.items():
        if any(keyword.lower() in lowered for keyword in keywords):
            flags.append(flag)
    return flags


def score_product(item: dict[str, Any]) -> tuple[int, str]:
    price = cents_to_float(item.get("sellPrice"))
    moq = item.get("startNum") or 999999
    try:
        moq_int = int(moq)
    except (TypeError, ValueError):
        moq_int = 999999

    flags = risk_flags(str(item.get("title") or ""))
    score = 0
    if price is not None:
        if price <= 3:
            score += 4
        elif price <= 10:
            score += 3
        elif price <= 30:
            score += 2
        else:
            score += 1
    if moq_int <= 5:
        score += 4
    elif moq_int <= 50:
        score += 3
    elif moq_int <= 200:
        score += 2
    else:
        score += 1
    score -= len(flags)

    if "ip" in flags:
        label = "Avoid"
    elif score >= 6:
        label = "A"
    elif score >= 4:
        label = "B"
    else:
        label = "C"
    return score, label


def normalize_product(item: dict[str, Any], source_query: str | None = None) -> dict[str, Any]:
    product_id = item.get("id")
    title = str(item.get("title") or "")
    score, label = score_product(item)
    return {
        "id": product_id,
        "title": title,
        "url": f"{WEB_ORIGIN}/product/detail/{product_id}.html",
        "price_min_yuan": cents_to_float(item.get("sellPrice")),
        "price_max_yuan": cents_to_float(item.get("maxPrice")),
        "price_display": price_display(item),
        "moq": item.get("startNum"),
        "shop_name": item.get("shopName"),
        "shop_credit": item.get("credit"),
        "picture": normalize_url(item.get("picture1")),
        "risk_flags": risk_flags(title),
        "score": score,
        "priority": label,
        "source_query": source_query,
        "raw": item,
    }


def build_keyword_query(keyword: str, constraints: str, count: int) -> str:
    return (
        f"请直接给我义乌购上具体可点击的商品链接：{keyword}。"
        f"要求：{constraints}。请推荐{count}个，按采购价、起购量、轻小件程度、跨境适配度排序，"
        "并说明每个商品的推荐理由和风险。"
    )


def sort_products(products: list[dict[str, Any]], mode: str) -> list[dict[str, Any]]:
    if mode == "ai":
        return products
    if mode == "price":
        return sorted(
            products,
            key=lambda p: (
                p.get("price_min_yuan") is None,
                numeric_sort_value(p.get("price_min_yuan")),
                numeric_sort_value(p.get("moq")),
            ),
        )
    return sorted(
        products,
        key=lambda p: (
            p.get("priority") == "Avoid",
            -numeric_sort_value(p.get("score"), default=0),
            numeric_sort_value(p.get("price_min_yuan")),
            numeric_sort_value(p.get("moq")),
        ),
    )


def build_markdown(chats: list[dict[str, Any]], products: list[dict[str, Any]], show_answers: bool = True) -> str:
    lines: list[str] = []
    for chat in chats:
        lines.append(f"session_id: {chat.get('session_id')}")
        if show_answers and chat.get("answer"):
            lines.append(chat["answer"].strip())
            lines.append("")

    if not products:
        lines.append("No product cards were returned. Retry with a shorter core keyword or a synonym.")
        return "\n".join(lines).rstrip() + "\n"

    lines.append("| Priority | ID | 商品 | 价格 | MOQ | 店铺 | 风险标记 | 链接 |")
    lines.append("|---|---|---|---:|---:|---|---|---|")
    for item in products:
        title = str(item.get("title") or "").replace("|", "/")
        shop = str(item.get("shop_name") or "").replace("|", "/")
        risks = ", ".join(item.get("risk_flags") or []) or "-"
        lines.append(
            f"| {item.get('priority')} | {item.get('id')} | {title} | {item.get('price_display')} | "
            f"{item.get('moq') or ''} | {shop} | {risks} | {item.get('url')} |"
        )
    return "\n".join(lines).rstrip() + "\n"


def print_markdown(chats: list[dict[str, Any]], products: list[dict[str, Any]], show_answers: bool = True) -> None:
    print(build_markdown(chats, products, show_answers=show_answers), end="")


def build_excel_image(image_url: str, max_size: int = 96) -> Any | None:
    if not image_url:
        return None
    try:
        from openpyxl.drawing.image import Image as ExcelImage
        from PIL import Image as PILImage
    except ImportError as exc:
        raise RuntimeError("Excel image export requires openpyxl and Pillow. Install them with: python3 -m pip install openpyxl Pillow") from exc

    try:
        data = request_bytes(normalize_url(image_url), {"User-Agent": BROWSER_UA, "Accept": "image/*"}, timeout=20)
        with PILImage.open(io.BytesIO(data)) as source:
            source.thumbnail((max_size, max_size))
            if source.mode not in {"RGB", "RGBA"}:
                source = source.convert("RGBA")
            image_data = io.BytesIO()
            source.save(image_data, format="PNG")
        image_data.seek(0)
        return ExcelImage(image_data)
    except Exception:
        return None


def write_excel_report(path: Path, products: list[dict[str, Any]], chats: list[dict[str, Any]]) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl. Install it with: python3 -m pip install openpyxl Pillow") from exc

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.freeze_panes = "A2"

    headers = [
        "Priority",
        "Preview",
        "ID",
        "Product",
        "Price",
        "MOQ",
        "Shop",
        "Risk Flags",
        "Score",
        "Link",
        "Source Query",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="FFF7ED")
    header_font = Font(bold=True, color="7C4A2B")
    thin_border = Border(bottom=Side(style="thin", color="E5E7EB"))
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    widths = {
        "A": 10,
        "B": 15,
        "C": 13,
        "D": 44,
        "E": 16,
        "F": 10,
        "G": 24,
        "H": 18,
        "I": 8,
        "J": 42,
        "K": 36,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    priority_fills = {
        "A": "DCFCE7",
        "B": "E0F2FE",
        "C": "FEF9C3",
        "Avoid": "FEE2E2",
    }
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A1:K{max(len(products) + 1, 2)}"

    for row_idx, item in enumerate(products, start=2):
        risks = ", ".join(item.get("risk_flags") or []) or "-"
        row = [
            item.get("priority"),
            "",
            item.get("id"),
            item.get("title"),
            item.get("price_display"),
            item.get("moq") or "",
            item.get("shop_name") or "",
            risks,
            item.get("score"),
            item.get("url"),
            item.get("source_query") or "",
        ]
        ws.append(row)
        ws.row_dimensions[row_idx].height = 76

        image = build_excel_image(item.get("picture") or "")
        if image:
            ws.add_image(image, f"B{row_idx}")
        else:
            ws.cell(row=row_idx, column=2).value = "No image"

        priority = str(item.get("priority") or "")
        if priority in priority_fills:
            ws.cell(row=row_idx, column=1).fill = PatternFill("solid", fgColor=priority_fills[priority])
        ws.cell(row=row_idx, column=10).hyperlink = item.get("url")
        ws.cell(row=row_idx, column=10).style = "Hyperlink"

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border

    meta = wb.create_sheet("Metadata")
    meta.append(["Generated At", dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    meta.append(["Products", len(products)])
    meta.append(["Session IDs", ", ".join(str(c.get("session_id")) for c in chats if c.get("session_id"))])
    meta.append([])
    meta.append(["AI Answers"])
    for idx, chat in enumerate(chats, start=1):
        meta.append([f"Answer {idx}", (chat.get("answer") or "").strip()])
    meta.column_dimensions["A"].width = 18
    meta.column_dimensions["B"].width = 100
    for row in meta.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].bestFit = True
    wb.save(path)
    return path


def html_escape(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def render_html_report(chats: list[dict[str, Any]], products: list[dict[str, Any]], title: str = "Yiwugo Sourcing Report") -> str:
    generated_at = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    session_ids = ", ".join(str(c.get("session_id")) for c in chats if c.get("session_id")) or "-"
    raw_answers = []
    rows: list[str] = []

    for item in products:
        risks = item.get("risk_flags") or []
        risk_html = "".join(f'<span class="risk">{html_escape(r)}</span>' for r in risks) or '<span class="muted">-</span>'
        image_url = item.get("picture") or ""
        if image_url:
            image_html = (
                f'<a href="{html_escape(item.get("url"))}" target="_blank" rel="noreferrer">'
                f'<img src="{html_escape(image_url)}" alt="{html_escape(item.get("title"))}" loading="lazy"></a>'
            )
        else:
            image_html = '<div class="no-image">No image</div>'
        rows.append(
            "<tr>"
            f'<td><span class="priority priority-{html_escape(item.get("priority"))}">{html_escape(item.get("priority"))}</span></td>'
            f'<td class="image-cell">{image_html}</td>'
            f'<td><a href="{html_escape(item.get("url"))}" target="_blank" rel="noreferrer">{html_escape(item.get("title"))}</a>'
            f'<div class="meta">ID: {html_escape(item.get("id"))}</div></td>'
            f'<td>{html_escape(item.get("price_display"))}</td>'
            f'<td>{html_escape(item.get("moq") or "")}</td>'
            f'<td>{html_escape(item.get("shop_name"))}<div class="meta">credit: {html_escape(item.get("shop_credit"))}</div></td>'
            f'<td>{risk_html}</td>'
            f'<td>{html_escape(item.get("score"))}</td>'
            "</tr>"
        )

    for idx, chat in enumerate(chats, start=1):
        answer = (chat.get("answer") or "").strip()
        if answer:
            raw_answers.append(f'<details><summary>AI answer {idx}</summary><pre>{html_escape(answer)}</pre></details>')

    if rows:
        table = (
            '<table><thead><tr>'
            '<th>Priority</th><th>主图</th><th>商品</th><th>价格</th><th>MOQ</th><th>店铺</th><th>风险标记</th><th>Score</th>'
            '</tr></thead><tbody>' + "\n".join(rows) + '</tbody></table>'
        )
    else:
        table = '<p class="empty">No product cards were returned. Retry with a shorter core keyword or a synonym.</p>'

    return f'''<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html_escape(title)}</title>
<style>
:root {{ --ink:#17202a; --muted:#6b7280; --line:#e5e7eb; --bg:#f7f4ee; --card:#fffdf8; --accent:#c45f2c; }}
* {{ box-sizing: border-box; }}
body {{ margin:0; color:var(--ink); background:linear-gradient(135deg,#f7f4ee,#eef6ef); font-family: ui-sans-serif, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif; }}
header {{ padding:32px 40px 24px; background:radial-gradient(circle at 20% 0%, #ffe3c6, transparent 36%), linear-gradient(90deg,#27211d,#6f3a1e); color:white; }}
h1 {{ margin:0 0 10px; font-size:32px; letter-spacing:.02em; }}
.summary {{ display:flex; flex-wrap:wrap; gap:12px; margin-top:14px; }}
.pill {{ padding:8px 12px; border:1px solid rgba(255,255,255,.35); border-radius:999px; background:rgba(255,255,255,.12); }}
main {{ padding:26px 40px 48px; }}
.card {{ background:var(--card); border:1px solid var(--line); border-radius:18px; box-shadow:0 16px 45px rgba(79,51,28,.08); overflow:hidden; }}
table {{ width:100%; border-collapse:collapse; }}
th, td {{ padding:14px 12px; border-bottom:1px solid var(--line); text-align:left; vertical-align:top; }}
th {{ font-size:12px; color:#7c4a2b; text-transform:uppercase; background:#fff7ed; position:sticky; top:0; z-index:1; }}
tr:hover {{ background:#fffaf1; }}
a {{ color:#8a3b16; text-decoration:none; }}
a:hover {{ text-decoration:underline; }}
.image-cell {{ width:116px; }}
img {{ width:96px; height:96px; object-fit:cover; border-radius:14px; border:1px solid var(--line); background:#f3f4f6; }}
.no-image {{ width:96px; height:96px; display:grid; place-items:center; color:var(--muted); background:#f3f4f6; border-radius:14px; font-size:12px; }}
.meta {{ margin-top:6px; font-size:12px; color:var(--muted); }}
.priority {{ display:inline-block; min-width:48px; text-align:center; padding:6px 10px; border-radius:999px; font-weight:700; font-size:12px; }}
.priority-A {{ background:#dcfce7; color:#166534; }}
.priority-B {{ background:#e0f2fe; color:#075985; }}
.priority-C {{ background:#fef9c3; color:#854d0e; }}
.priority-Avoid {{ background:#fee2e2; color:#991b1b; }}
.risk {{ display:inline-block; margin:0 5px 5px 0; padding:5px 8px; border-radius:999px; background:#f3f4f6; color:#374151; font-size:12px; }}
.muted, .empty {{ color:var(--muted); }}
details {{ margin-top:20px; background:#fff; border:1px solid var(--line); border-radius:14px; padding:12px 14px; }}
summary {{ cursor:pointer; font-weight:700; }}
pre {{ white-space:pre-wrap; word-break:break-word; color:#374151; }}
@media (max-width: 820px) {{ header, main {{ padding-left:18px; padding-right:18px; }} table {{ font-size:13px; }} th:nth-child(8), td:nth-child(8), th:nth-child(6), td:nth-child(6) {{ display:none; }} }}
</style>
</head>
<body>
<header>
<h1>{html_escape(title)}</h1>
<div>Generated at {html_escape(generated_at)}</div>
<div class="summary"><span class="pill">Queries: {len(chats)}</span><span class="pill">Products: {len(products)}</span><span class="pill">Sessions: {html_escape(session_ids)}</span></div>
</header>
<main>
<section class="card">{table}</section>
<section>{''.join(raw_answers)}</section>
</main>
</body>
</html>
'''


def default_report_prefix() -> Path:
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path.cwd() / f"yiwugo_sourcing_report_{stamp}"


def write_reports(prefix: Path, products: list[dict[str, Any]], chats: list[dict[str, Any]], html_report: str, markdown: str | None = None) -> tuple[Path, Path, Path | None]:
    prefix.parent.mkdir(parents=True, exist_ok=True)
    xlsx_path = prefix.with_suffix(".xlsx")
    html_path = prefix.with_suffix(".html")
    write_excel_report(xlsx_path, products, chats)
    html_path.write_text(html_report, encoding="utf-8")
    md_path = None
    if markdown is not None:
        md_path = prefix.with_suffix(".md")
        md_path.write_text(markdown, encoding="utf-8")
    return xlsx_path, html_path, md_path


def run_queries(messages: list[str], client_id: str | None, session_id: str | None, timeout: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    chats: list[dict[str, Any]] = []
    products: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    current_session = session_id
    shared_client = client_id or str(uuid.uuid4())

    for message in messages:
        chat = stream_chat(message, client_id=shared_client, session_id=current_session, timeout=timeout)
        if session_id:
            current_session = chat.get("session_id") or current_session
        chats.append(chat)
        for raw in fetch_products(chat.get("product_ids", [])):
            product_id = str(raw.get("id"))
            if product_id in seen_ids:
                continue
            seen_ids.add(product_id)
            products.append(normalize_product(raw, source_query=message))
    return chats, products


def main() -> int:
    parser = argparse.ArgumentParser(description="Query Yiwugo AI without Chrome")
    parser.add_argument("message", nargs="?", help="Chinese sourcing query to send to Yiwugo AI")
    parser.add_argument("--keyword", action="append", help="Core product keyword. Repeat for fallback/synonym searches.")
    parser.add_argument("--constraints", default="轻小件、采购价低、适合跨境电商", help="Natural-language constraints for --keyword queries")
    parser.add_argument("--count", type=int, default=8, help="Requested number of products per keyword query")
    parser.add_argument("--max-products", type=int, default=20, help="Maximum products to print after de-duplication")
    parser.add_argument("--sort", choices=["score", "price", "ai"], default="score", help="Sort products by heuristic score, price, or AI order")
    parser.add_argument("--client-id", help="Stable UUID-like userAgent value for continued conversations")
    parser.add_argument("--session-id", help="Session ID returned by a previous call")
    parser.add_argument("--timeout", type=int, default=120, help="Per-query timeout in seconds")
    parser.add_argument("--no-answer", action="store_true", help="Hide the raw AI answer in terminal preview output")
    parser.add_argument("--json", action="store_true", help="Print raw JSON instead of terminal preview/report files")
    parser.add_argument("--markdown", action="store_true", help="Also write a Markdown report beside the default Excel and HTML reports")
    parser.add_argument("--out-prefix", help="Report output path prefix without extension. Default writes timestamped .xlsx and .html files in cwd.")
    parser.add_argument("--no-report-files", action="store_true", help="Do not write default Excel and HTML report files")
    parser.add_argument("--report-title", default="Yiwugo Sourcing Report", help="Title for the generated HTML report")
    args = parser.parse_args()

    messages: list[str] = []
    if args.message:
        messages.append(args.message)
    for keyword in args.keyword or []:
        messages.append(build_keyword_query(keyword, args.constraints, args.count))
    if not messages:
        parser.error("provide a message or at least one --keyword")

    chats, products = run_queries(messages, args.client_id, args.session_id, args.timeout)
    products = sort_products(products, args.sort)[: args.max_products]
    result = {"chats": chats, "products": products}
    if args.json:
        json.dump(result, sys.stdout, ensure_ascii=False, indent=2)
        print()
    else:
        markdown = build_markdown(chats, products, show_answers=not args.no_answer)
        print(markdown, end="")
        if not args.no_report_files:
            prefix = Path(args.out_prefix).expanduser() if args.out_prefix else default_report_prefix()
            html_report = render_html_report(chats, products, title=args.report_title)
            xlsx_path, html_path, md_path = write_reports(prefix, products, chats, html_report, markdown=markdown if args.markdown else None)
            print(f"\nExcel report: {xlsx_path}")
            print(f"HTML report: {html_path}")
            if md_path:
                print(f"Markdown report: {md_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
